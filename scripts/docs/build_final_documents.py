from __future__ import annotations

import csv
import datetime as dt
import shutil
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "entregas"
TMP_IMG_DIR = ROOT / "tmp" / "doc_assets"

TODAY = dt.date.today().strftime("%d/%m/%Y")


COLORS = {
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "navy": "0B2545",
    "muted": "666666",
    "border": "B7C3D0",
    "table_header": "F2F4F7",
    "callout": "F4F6F9",
    "success": "EAF4EA",
    "caution": "FFF8E5",
}


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def count_csv_records(path: Path) -> int:
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8", newline="") as handle:
        return max(sum(1 for _ in handle) - 1, 0)


def count_files(path: Path) -> int:
    if not path.exists():
        return 0
    return sum(1 for item in path.rglob("*") if item.is_file())


def pct(value: str | float, digits: int = 1) -> str:
    try:
        return f"{float(value) * 100:.{digits}f}%"
    except (TypeError, ValueError):
        return "-"


def num(value: str | float, digits: int = 3) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return "-"


def ms(value: str | float) -> str:
    try:
        return f"{float(value):.0f} ms"
    except (TypeError, ValueError):
        return "-"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9DEE8", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: Sequence[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
                set_cell_margins(row.cells[idx])
                set_cell_border(row.cells[idx])
                row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def set_run_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document, preset: str) -> None:
    styles = doc.styles
    base = styles["Normal"]
    base.font.name = "Calibri"
    base._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    base._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    base.font.size = Pt(11)
    base.paragraph_format.space_after = Pt(6 if preset == "brief" else 8)
    base.paragraph_format.line_spacing = 1.10 if preset == "brief" else 1.25
    base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name, size, color, before, after in [
        ("Title", 23, "000000", 0, 4),
        ("Subtitle", 13, "555555", 0, 12),
        ("Heading 1", 16, COLORS["blue"], 16 if preset == "brief" else 18, 8),
        ("Heading 2", 13, COLORS["blue"], 12, 6),
        ("Heading 3", 12, COLORS["dark_blue"], 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name.startswith("Heading") or style_name == "Title":
            style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = style_name.startswith("Heading")

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167 if preset == "brief" else 1.208


def configure_page(doc: Document, title: str, subtitle: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run(title)
    set_run_font(run, size=9, color=COLORS["muted"], bold=True)
    header.add_run(" | ")
    run = header.add_run(subtitle)
    set_run_font(run, size=9, color=COLORS["muted"])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Documento gerado em " + TODAY)
    set_run_font(run, size=8.5, color=COLORS["muted"])


def add_title_block(doc: Document, title: str, subtitle: str, metadata: Sequence[tuple[str, str]]) -> None:
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(8)
    p.add_run(title)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run(subtitle)

    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_width(table, [1.45, 5.05])
    mark_header_row(table)
    for idx, (label, value) in enumerate(metadata):
        left, right = table.rows[idx].cells
        set_cell_shading(left, COLORS["table_header"])
        left.paragraphs[0].add_run(label).bold = True
        right.paragraphs[0].add_run(value)


def add_callout(doc: Document, title: str, body: str, fill: str = "F4F6F9") -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    mark_header_row(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    p.add_run(body)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    mark_header_row(table)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(hdr[idx], COLORS["table_header"])
        paragraph = hdr[idx].paragraphs[0]
        run = paragraph.add_run(header)
        run.bold = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if widths[idx] <= 1.2 else WD_ALIGN_PARAGRAPH.LEFT
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            paragraph = cells[idx].paragraphs[0]
            paragraph.add_run(value)
            if widths[idx] <= 1.2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_width(table, widths)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=COLORS["muted"], italic=True)


def make_image_copy(source: Path, max_width_px: int = 1600) -> Path | None:
    if not source.exists():
        return None
    TMP_IMG_DIR.mkdir(parents=True, exist_ok=True)
    target = TMP_IMG_DIR / (source.stem + "_doc.jpg")
    with Image.open(source) as image:
        image = image.convert("RGB")
        if image.width > max_width_px:
            height = int(image.height * (max_width_px / image.width))
            image = image.resize((max_width_px, height), Image.LANCZOS)
        image.save(target, "JPEG", quality=86, optimize=True)
    return target


def add_figure(doc: Document, image: Path, caption: str, width: float = 6.1) -> None:
    processed = make_image_copy(image)
    if processed is None:
        add_callout(doc, "Figura indisponivel", f"O arquivo {image} nao foi encontrado.", fill=COLORS["caution"])
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(processed), width=Inches(width))
    inline._inline.docPr.set("title", "Figura")
    inline._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def load_metrics() -> dict[str, dict[str, str]]:
    face = read_rows(ROOT / "outputs" / "reports" / "gallagher_face" / "face_metrics.csv")[0]
    global_person = read_rows(ROOT / "outputs" / "reports" / "gallagher_global_person" / "global_metrics.csv")[0]
    holidays = read_rows(ROOT / "outputs" / "reports" / "holidays_global" / "global_metrics.csv")[0]
    fusion_rows = read_rows(ROOT / "outputs" / "reports" / "gallagher_fusion" / "fusion_metrics.csv")
    best_fusion = max(fusion_rows, key=lambda row: float(row["mAP"]))
    return {
        "face": face,
        "global_person": global_person,
        "holidays": holidays,
        "fusion": best_fusion,
    }


def metric_rows(metrics: dict[str, dict[str, str]]) -> list[list[str]]:
    return [
        ["Gallagher", "Face-only", metrics["face"]["num_queries"], pct(metrics["face"]["precision_at_5"]), pct(metrics["face"]["precision_at_10"]), pct(metrics["face"]["recall_at_10"]), num(metrics["face"]["mAP"]), ms(metrics["face"]["query_time_ms"])],
        ["Gallagher", "Global ResNet50", metrics["global_person"]["num_queries"], pct(metrics["global_person"]["precision_at_5"]), pct(metrics["global_person"]["precision_at_10"]), pct(metrics["global_person"]["recall_at_10"]), num(metrics["global_person"]["mAP"]), ms(metrics["global_person"]["query_time_ms"])],
        ["Gallagher", "Fusao 0.9/0.1", metrics["fusion"]["num_queries"], pct(metrics["fusion"]["precision_at_5"]), pct(metrics["fusion"]["precision_at_10"]), pct(metrics["fusion"]["recall_at_10"]), num(metrics["fusion"]["mAP"]), ms(metrics["fusion"]["query_time_ms"])],
        ["INRIA Holidays", "Global ResNet50", metrics["holidays"]["num_queries"], pct(metrics["holidays"]["precision_at_5"]), pct(metrics["holidays"]["precision_at_10"]), pct(metrics["holidays"]["recall_at_10"]), num(metrics["holidays"]["mAP"]), ms(metrics["holidays"]["query_time_ms"])],
    ]


def add_references(doc: Document, compact: bool = False, heading: str = "Referencias") -> None:
    doc.add_heading(heading, level=1)
    refs = [
        "Smeulders, A. W. M. et al. Content-Based Image Retrieval at the End of the Early Years. IEEE TPAMI, 2000. DOI: 10.1109/34.895972.",
        "Datta, R. et al. Image Retrieval: Ideas, Influences, and Trends of the New Age. ACM Computing Surveys, 2008. DOI: 10.1145/1348246.1348248.",
        "Dubey, S. R. A Decade Survey of Content Based Image Retrieval using Deep Learning. IEEE TCSVT, 2021. DOI: 10.1109/TCSVT.2021.3080920.",
        "Razavian, A. S. et al. CNN Features Off-the-Shelf: An Astounding Baseline for Recognition. CVPR Workshops, 2014.",
        "Babenko, A. et al. Neural Codes for Image Retrieval. ECCV, 2014.",
        "He, K. et al. Deep Residual Learning for Image Recognition. CVPR, 2016. DOI: 10.1109/CVPR.2016.90.",
        "Schroff, F.; Kalenichenko, D.; Philbin, J. FaceNet: A Unified Embedding for Face Recognition and Clustering. CVPR, 2015.",
        "Wang, H. et al. CosFace: Large Margin Cosine Loss for Deep Face Recognition. CVPR, 2018.",
        "Deng, J. et al. ArcFace: Additive Angular Margin Loss for Deep Face Recognition. CVPR, 2019.",
        "Huang, G. B. et al. Labeled Faces in the Wild: A Database for Studying Face Recognition in Unconstrained Environments. 2007.",
        "Jegou, H.; Douze, M.; Schmid, C. Hamming Embedding and Weak Geometric Consistency for Large Scale Image Search. ECCV, 2008.",
        "Gallagher, A. C.; Das, M.; Loui, A. C. User-Assisted People Search in Consumer Image Collections. ICME, 2007.",
        "Zhang, N. et al. Beyond Frontal Faces: Improving Person Recognition Using Multiple Cues. CVPR, 2015.",
        "Oh, S. J. et al. Person Recognition in Personal Photo Collections. ICCV, 2015.",
        "Li, H. et al. A Multi-Level Contextual Model for Person Recognition in Photo Albums. CVPR, 2016.",
        "Snoek, C. G. M.; Worring, M.; Smeulders, A. W. M. Early versus Late Fusion in Semantic Video Analysis. ACM Multimedia, 2005.",
        "Manning, C. D.; Raghavan, P.; Schutze, H. Introduction to Information Retrieval. Cambridge University Press, 2008.",
        "Buolamwini, J.; Gebru, T. Gender Shades: Intersectional Accuracy Disparities in Commercial Gender Classification. PMLR, 2018.",
        "NIST. Face Recognition Vendor Test Part 3: Demographic Effects. NIST IR 8280, 2019.",
    ]
    if compact:
        refs = refs[:12]
    for idx, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"[{idx}] {ref}")


def build_report(metrics: dict[str, dict[str, str]]) -> Path:
    doc = Document()
    configure_styles(doc, "brief")
    configure_page(doc, "Relatorio da Meta 2", "Recuperacao fotografica face + CBIR global")
    doc.core_properties.title = "Relatorio da Meta 2 - Recuperacao Fotografica Face + CBIR Global"
    doc.core_properties.author = "Artur Duarte"

    for text in [
        "Instituto Federal de Educacao, Ciencia e Tecnologia Goiano - Campus Urutai",
        "Discente: Artur Duarte Monteiro",
        "Orientador: Gabriel da Silva Vieira",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=11, bold=text.startswith("Instituto"))

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Relatorio Meta 2")
    set_run_font(r, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Recuperacao de Imagens Fotograficas por Conteudo e Reconhecimento Facial em Colecoes de Grande Escala")
    set_run_font(r, size=13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Busca facial multi-rosto, CBIR global com ResNet50 e fusao tardia de scores")
    set_run_font(r, size=11, color=COLORS["muted"])

    for _ in range(9):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Urutai - GO, 2026")
    set_run_font(r, size=11)

    doc.add_page_break()

    doc.add_heading("1 Introducao", level=1)
    doc.add_paragraph(
        "Este relatorio descreve as atividades executadas na Meta 2 do projeto, com foco na consolidacao "
        "de um sistema reprodutivel de recuperacao fotografica por conteudo. Nesta etapa, o projeto deixou "
        "de ser apenas um conjunto de scripts exploratorios e passou a possuir modulos separados para busca "
        "facial, busca global da imagem inteira e fusao tardia de scores."
    )
    doc.add_paragraph(
        "O objetivo tecnico da Meta 2 foi implementar e avaliar tres estrategias de recuperacao: busca por "
        "pessoa com embeddings faciais, busca por similaridade visual com descritores globais ResNet50 e "
        "combinacao ponderada entre as duas trilhas. O resultado esperado e uma base concreta para relatorio, "
        "TCC ou artigo curto, com metricas objetivas e exemplos visuais."
    )

    doc.add_heading("2 Metodologia", level=1)
    doc.add_heading("2.1 Organizacao do sistema", level=2)
    doc.add_paragraph(
        "A implementacao foi organizada em diretorios por responsabilidade. A pasta scripts/face contem os "
        "programas de indexacao, busca e avaliacao facial. A pasta scripts/global contem os scripts de CBIR "
        "global com ResNet50. A pasta scripts/fusion contem a combinacao dos rankings, enquanto scripts/app "
        "armazena a interface minima em Streamlit."
    )
    doc.add_paragraph(
        "Essa separacao foi adotada para permitir a avaliacao isolada de cada modulo antes da fusao. Assim, "
        "quando um resultado melhora ou piora, e possivel identificar se a causa esta no modulo facial, no "
        "modulo global ou na regra de combinacao."
    )

    doc.add_heading("2.2 Indice facial multi-rosto", level=2)
    doc.add_paragraph(
        "O indice facial foi construido considerando que uma mesma fotografia pode conter varias pessoas. "
        "Portanto, a unidade indexada nao e a imagem inteira, mas cada rosto detectado dentro da imagem. "
        "Cada face recebe um embedding proprio e metadados que apontam para a fotografia original, incluindo "
        "identificador da imagem, posicao do rosto e linha correspondente no arquivo de embeddings."
    )

    doc.add_heading("2.3 CBIR global com ResNet50", level=2)
    doc.add_paragraph(
        "Para a busca global, cada fotografia inteira foi processada por uma ResNet50 pre-treinada. O vetor "
        "extraido representa o conteudo visual geral da imagem, incluindo cena, composicao, objetos e padroes "
        "visuais. Esse modulo foi mantido separado do facial porque ele responde a uma pergunta diferente: "
        "quais imagens sao visualmente semelhantes a imagem de consulta."
    )

    doc.add_heading("2.4 Fusao tardia de scores", level=2)
    doc.add_paragraph(
        "A fusao foi implementada como late fusion, ou fusao tardia. Em vez de concatenar embeddings faciais "
        "e globais, que possuem significados diferentes, o sistema combina os scores finais de similaridade. "
        "Foram testados os pesos 0.9/0.1, 0.7/0.3 e 0.5/0.5 para face/global."
    )
    add_table(
        doc,
        ["Modulo", "Entrada", "Saida gerada", "Finalidade"],
        [
            ["Face", "Foto com uma ou mais pessoas", "Embedding facial de 512 dimensoes por rosto", "Encontrar a mesma pessoa em outras fotos"],
            ["Global", "Imagem inteira", "Embedding global ResNet50 de 2048 dimensoes", "Encontrar imagens visualmente semelhantes"],
            ["Fusao", "Scores facial e global", "Score final ponderado", "Testar se contexto visual melhora o ranking"],
            ["Interface", "Upload do usuario", "Galeria ranqueada", "Demonstrar a operacao do sistema"],
        ],
        [1.05, 1.55, 2.0, 1.9],
    )

    doc.add_heading("3 Bases de dados e organizacao experimental", level=1)
    doc.add_paragraph(
        "A avaliacao foi dividida por tarefa para evitar conclusoes incorretas. A base LFW foi usada como "
        "validacao inicial de reconhecimento facial. A base Gallagher foi usada para busca por pessoas em "
        "fotografias de grupo e para avaliacao da fusao. A base INRIA Holidays foi usada para avaliar a busca "
        "global por similaridade visual de imagem inteira."
    )
    add_table(
        doc,
        ["Base/indice", "Escopo verificado", "Uso no projeto"],
        [
            ["LFW", "13.233 imagens processadas e 16.058 embeddings faciais", "Validacao inicial de reconhecimento facial"],
            ["Gallagher", "589 imagens, 931 faces anotadas e 32 identidades", "Avaliacao de busca por pessoa em fotos de grupo"],
            ["INRIA Holidays", "1.491 imagens, 500 consultas e 991 relacoes de relevancia", "Avaliacao de CBIR global por imagem inteira"],
            ["Indice facial Gallagher", f"{count_csv_records(ROOT / 'outputs/face_index_gallagher/face_metadata.csv')} embeddings faciais", "Busca facial e fusao"],
            ["Indice global Holidays", f"{count_csv_records(ROOT / 'outputs/global_index_holidays/global_metadata.csv')} embeddings globais", "Busca global ResNet50"],
        ],
        [1.55, 2.25, 2.7],
    )

    doc.add_heading("4 Implementacao dos modulos", level=1)
    doc.add_paragraph(
        "A Meta 2 gerou scripts novos e reprodutiveis para indexacao, busca, avaliacao e demonstracao. "
        "Os scripts principais aceitam parametros por linha de comando, como diretorio de entrada, diretorio "
        "de saida, dispositivo de execucao, Top-K e limiar de similaridade."
    )
    add_table(
        doc,
        ["Script", "Funcao", "Saida principal"],
        [
            ["scripts/face/01_index_faces.py", "Detecta todos os rostos e gera embeddings faciais.", "outputs/face_index*/face_embeddings.npy"],
            ["scripts/face/02_search_person.py", "Busca uma pessoa a partir de uma imagem de consulta.", "topk_results.csv e exemplos visuais"],
            ["scripts/face/03_evaluate_face.py", "Avalia Precision@K, Recall@K, mAP e tempo.", "outputs/reports/*face*"],
            ["scripts/global/01_index_global_resnet.py", "Extrai embeddings globais ResNet50.", "global_embeddings.npy"],
            ["scripts/global/02_search_global.py", "Busca imagens semelhantes pela imagem inteira.", "global_topk_results.csv"],
            ["scripts/global/03_evaluate_global.py", "Avalia o CBIR global.", "outputs/reports/*global*"],
            ["scripts/fusion/02_evaluate_fusion.py", "Compara face, global e fusao.", "fusion_metrics.csv"],
            ["scripts/app/streamlit_app.py", "Demonstra busca por pessoa, global e fusao.", "Interface local em Streamlit"],
        ],
        [2.15, 2.45, 1.9],
    )

    doc.add_heading("5 Avaliacao experimental", level=1)
    doc.add_paragraph(
        "As metricas abaixo foram calculadas com Precision@K, Recall@K, mAP e tempo medio de consulta. "
        "No Gallagher, a avaliacao mede recuperacao de pessoas. No INRIA Holidays, a avaliacao mede "
        "similaridade visual entre imagens completas."
    )
    add_table(
        doc,
        ["Base", "Metodo", "Queries", "P@5", "P@10", "R@10", "mAP", "Tempo"],
        metric_rows(metrics),
        [1.15, 1.45, 0.65, 0.65, 0.65, 0.65, 0.75, 0.75],
    )

    doc.add_heading("5.1 Busca facial no Gallagher", level=2)
    doc.add_paragraph(
        "O resultado facial no Gallagher confirma que embeddings faciais sao adequados quando a pergunta do "
        f"usuario e encontrar uma pessoa especifica. O metodo face-only obteve P@5 de {pct(metrics['face']['precision_at_5'])}, "
        f"P@10 de {pct(metrics['face']['precision_at_10'])} e mAP de {num(metrics['face']['mAP'])}. "
        "Como o sistema retorna fotos completas, ele atende ao objetivo pratico de localizar fotografias onde "
        "a pessoa aparece, e nao apenas recortes de rosto."
    )

    doc.add_heading("5.2 CBIR global no INRIA Holidays", level=2)
    doc.add_paragraph(
        f"No INRIA Holidays, o modulo global com ResNet50 obteve mAP de {num(metrics['holidays']['mAP'])}. "
        "Esse resultado indica que a extracao de descritores globais esta coerente com a tarefa de recuperar "
        "imagens visualmente semelhantes. No Gallagher, quando a relevancia foi definida por identidade, o "
        "mesmo modulo teve desempenho baixo, pois captura cena, fundo, objetos e composicao visual, nao identidade."
    )

    doc.add_heading("5.3 Fusao face + global", level=2)
    doc.add_paragraph(
        "A fusao 0.9/0.1 foi a melhor entre os pesos testados, mas ainda ficou abaixo da busca facial isolada. "
        "Isso indica que adicionar contexto global pode introduzir ruido quando o criterio de relevancia e a "
        "identidade da pessoa. Para artigo, esse resultado deve ser apresentado como uma evidencia experimental: "
        "a fusao simples nao melhora automaticamente todos os cenarios."
    )

    doc.add_heading("6 Interface e demonstracao", level=1)
    doc.add_paragraph(
        "Foi criada uma interface minima em Streamlit para demonstrar o funcionamento do sistema. A interface "
        "permite enviar uma imagem de consulta, escolher o modo de busca, ajustar Top-K e limiar facial, e "
        "visualizar os resultados em uma galeria ranqueada. A interface nao e o foco cientifico da Meta 2, mas "
        "serve como prova de conceito e facilita a apresentacao do sistema."
    )

    doc.add_heading("7 Evidencias e Reprodutibilidade", level=1)
    doc.add_paragraph(
        "Os experimentos foram executados com ambiente Conda dedicado e suporte a GPU CUDA para os modulos "
        "compativeis. Os resultados quantitativos foram salvos em CSV, e os exemplos visuais foram salvos em "
        "outputs/reports. Os dados brutos, embeddings e outputs permanecem fora do Git para manter o repositorio leve."
    )
    add_table(
        doc,
        ["Evidencia", "Caminho"],
        [
            ["Metricas faciais", "outputs/reports/gallagher_face/face_metrics.csv"],
            ["Metricas globais Gallagher", "outputs/reports/gallagher_global_person/global_metrics.csv"],
            ["Metricas globais Holidays", "outputs/reports/holidays_global/global_metrics.csv"],
            ["Metricas de fusao", "outputs/reports/gallagher_fusion/fusion_metrics.csv"],
            ["Tabela comparativa", "outputs/reports/comparison/gallagher_face_global_fusion_comparison.csv"],
            ["Interface", "scripts/app/streamlit_app.py"],
        ],
        [2.1, 4.4],
    )
    add_numbered(
        doc,
        [
            "Indexar rostos: python scripts/face/01_index_faces.py --input-dir <pasta> --output-dir <saida> --device cuda",
            "Indexar imagens globais: python scripts/global/01_index_global_resnet.py --input-dir <pasta> --output-dir <saida> --device cuda",
            "Avaliar face, global e fusao com os scripts 03_evaluate_face.py, 03_evaluate_global.py e 02_evaluate_fusion.py.",
            "Executar a interface: python -m streamlit run scripts/app/streamlit_app.py --server.port 8501",
        ],
    )

    doc.add_heading("8 Resultados e Discussao", level=1)
    doc.add_paragraph(
        "Os resultados mostram que a busca facial e a estrategia mais adequada quando o objetivo e recuperar "
        "fotografias de uma pessoa especifica. Ja a busca global e mais apropriada quando o objetivo e encontrar "
        "imagens parecidas pela cena ou composicao visual. A fusao simples nao superou a busca facial isolada no "
        "Gallagher, mas esse resultado e relevante: ele demonstra que combinar sinais heterogeneos sem calibracao "
        "pode prejudicar o ranking."
    )
    doc.add_paragraph(
        "Para fins de artigo, a principal contribuicao nao e afirmar que a fusao sempre melhora, mas mostrar uma "
        "avaliacao comparativa clara entre face-only, global-only e fusao tardia. A conclusao tecnica e que os "
        "modulos devem permanecer separados na interface e que a fusao deve ser tratada como recurso experimental."
    )

    doc.add_heading("8.1 Exemplos visuais", level=2)
    add_figure(
        doc,
        ROOT / "outputs" / "reports" / "gallagher_face" / "visual_examples" / "face_examples.png",
        "Figura 1 - Exemplo de busca facial no Gallagher: a consulta por uma pessoa retorna fotos completas do acervo.",
    )
    add_figure(
        doc,
        ROOT / "outputs" / "reports" / "holidays_global" / "visual_examples" / "global_examples.png",
        "Figura 2 - Exemplo de busca global no INRIA Holidays: o ranking e baseado na similaridade visual da imagem inteira.",
    )

    doc.add_heading("8.2 Criterios de aceite", level=2)
    add_table(
        doc,
        ["Criterio", "Status", "Evidencia"],
        [
            ["Repositorio leve", "Atendido", "Bases, embeddings, modelos e outputs estao fora do Git."],
            ["Scripts por CLI", "Atendido", "Scripts novos aceitam argumentos como input-dir, output-dir, topk, threshold e device."],
            ["Indices reproduziveis", "Atendido", "Indices facial e global sao salvos em NPY + CSV de metadados."],
            ["Metricas em CSV", "Atendido", "Relatorios gerados em outputs/reports com Precision@K, Recall@K, mAP e tempo."],
            ["Tabela face/global/fusao", "Atendido", "Comparacao consolidada em outputs/reports/comparison."],
            ["Interface minima", "Atendido", "Streamlit executado localmente com busca por pessoa, global e fusao."],
            ["Escala de milhoes", "Futuro", "V1 usa NumPy; FAISS fica para etapa posterior."],
        ],
        [2.2, 1.0, 3.3],
    )

    doc.add_heading("9 Conclusao", level=1)
    doc.add_paragraph(
        "A Meta 2 pode ser considerada concluida em nivel tecnico e experimental para uma v1. O sistema possui "
        "modulos separados de busca facial e CBIR global, avaliacao quantitativa, exemplos visuais e interface "
        "minima. O material produzido ja sustenta a escrita de um artigo curto ou relatorio tecnico, desde que "
        "as limitacoes sejam apresentadas de forma transparente."
    )

    doc.add_heading("10 Proximos Passos", level=1)
    add_bullets(
        doc,
        [
            "Ampliar a avaliacao facial com mais consultas e, se possivel, outro acervo de eventos.",
            "Permitir selecao manual do rosto de consulta quando uma imagem possuir mais de uma face.",
            "Testar CLIP como descritor global complementar ao baseline ResNet50.",
            "Adicionar FAISS para acelerar a busca em acervos maiores.",
            "Investigar normalizacao de scores e pesos aprendidos para melhorar a fusao.",
            "Documentar cuidados de privacidade, vies e armazenamento seguro de embeddings faciais.",
            "Converter o rascunho tecnico em artigo no formato exigido pelo evento ou periodico escolhido.",
        ],
    )

    add_references(doc, compact=True, heading="11 Referencias")

    output = OUT_DIR / "Relatorio_Meta2_CIBIR.docx"
    doc.save(output)
    return output


def build_article(metrics: dict[str, dict[str, str]]) -> Path:
    doc = Document()
    configure_styles(doc, "proposal")
    configure_page(doc, "Rascunho de artigo", "Face + CBIR global")
    doc.core_properties.title = "Rascunho de artigo - Recuperacao Fotografica Face + CBIR Global"
    doc.core_properties.author = "Artur Duarte"

    add_title_block(
        doc,
        "Recuperacao Fotografica por Reconhecimento Facial e CBIR Global com Fusao Tardia de Scores",
        "Rascunho de artigo tecnico-cientifico",
        [
            ("Autor", "Artur Duarte"),
            ("Linha", "Recuperacao de imagens por conteudo, reconhecimento facial e avaliacao experimental"),
            ("Versao", "Rascunho inicial para orientacao e submissao futura"),
            ("Data", TODAY),
        ],
    )

    doc.add_heading("Resumo", level=1)
    doc.add_paragraph(
        "O crescimento de acervos fotograficos digitais torna dificil localizar manualmente imagens de pessoas "
        "ou cenas especificas. Este trabalho apresenta uma versao inicial de um sistema de recuperacao fotografica "
        "que combina duas formas de representacao visual: embeddings faciais para busca por pessoa e descritores "
        "globais extraidos com ResNet50 para busca por similaridade da imagem inteira. A arquitetura indexa "
        "multiplos rostos por fotografia, retorna a imagem completa ao usuario e permite comparar busca facial, "
        "busca global e fusao tardia de scores por similaridade cosseno. A avaliacao foi conduzida com LFW como "
        "validacao facial inicial, Gallagher para busca por pessoas em fotos pessoais e INRIA Holidays para CBIR "
        "global. Nos experimentos, a busca facial obteve P@5 de "
        f"{pct(metrics['face']['precision_at_5'])} e mAP de {num(metrics['face']['mAP'])} no Gallagher; a busca global "
        f"obteve mAP de {num(metrics['holidays']['mAP'])} no INRIA Holidays; e a fusao simples 0.9/0.1 nao superou "
        "a busca facial isolada na tarefa centrada em identidade. Os resultados indicam que face e descritor global "
        "sao complementares em intencao de uso, mas a combinacao linear simples deve ser aplicada com cautela."
    )
    doc.add_paragraph("Palavras-chave: CBIR; reconhecimento facial; embeddings; ResNet50; fusao tardia; recuperacao de imagens.")

    doc.add_heading("1. Introducao", level=1)
    doc.add_paragraph(
        "Acervos fotograficos pessoais e institucionais crescem rapidamente em eventos, cerimonias, atividades "
        "academicas, albuns familiares e colecoes historicas. Em muitos casos, o usuario nao procura uma imagem "
        "pelo nome do arquivo, mas pelo conteudo: uma pessoa, um grupo, uma cena ou uma composicao visual semelhante. "
        "Essa motivacao aproxima o projeto da area de Content-Based Image Retrieval, em que a recuperacao e feita "
        "a partir das caracteristicas visuais extraidas da propria imagem."
    )
    doc.add_paragraph(
        "A proposta deste trabalho e avaliar uma arquitetura modular que une busca por rosto e busca global por "
        "imagem inteira. A busca por rosto responde a pergunta 'em quais fotos esta esta pessoa?'. A busca global "
        "responde a pergunta 'quais fotos sao visualmente parecidas com esta imagem?'. A fusao tardia busca verificar "
        "se a combinacao de identidade facial e contexto visual melhora o ranking final."
    )
    add_callout(
        doc,
        "Contribuicao esperada",
        "O trabalho nao se limita a demonstrar uma interface. Ele compara experimentalmente tres estrategias de "
        "recuperacao: face-only, global-only e fusao face + global, retornando fotografias completas e usando "
        "metricas de ranking.",
        fill=COLORS["callout"],
    )

    doc.add_heading("2. Trabalhos relacionados", level=1)
    doc.add_heading("2.1 CBIR e descritores globais", level=2)
    doc.add_paragraph(
        "A literatura classica de CBIR discute o uso de cor, textura, forma e outras caracteristicas visuais para "
        "recuperar imagens por conteudo, alem do problema conhecido como semantic gap, isto e, a distancia entre "
        "descritores numericos e conceitos percebidos por humanos [1, 2]. Com o avanco de deep learning, redes "
        "convolucionais pre-treinadas passaram a ser usadas como extratores genericos de caracteristicas, mesmo "
        "sem treinamento especifico para o novo dominio [3, 4, 5]. Neste projeto, a ResNet50 e usada como baseline "
        "global por ser uma arquitetura consolidada e simples de reproduzir [6]."
    )
    doc.add_heading("2.2 Reconhecimento facial por embeddings", level=2)
    doc.add_paragraph(
        "Sistemas modernos de reconhecimento facial usualmente transformam cada face em um vetor numerico, ou "
        "embedding, no qual faces da mesma pessoa tendem a ficar proximas. FaceNet, CosFace e ArcFace consolidaram "
        "essa linha ao introduzir perdas e espacos de representacao adequados a verificacao e identificacao facial "
        "[7, 8, 9]. A vantagem para este projeto e que a pessoa consultada nao precisa ser uma classe treinada: basta "
        "extrair seu embedding e compara-lo com o indice."
    )
    doc.add_heading("2.3 Pessoas em albuns fotograficos", level=2)
    doc.add_paragraph(
        "A busca de pessoas em albuns pessoais e mais dificil que a verificacao facial controlada, pois as imagens "
        "incluem oclusoes, baixa resolucao, variacao de pose, roupas, fundo, iluminacao e coocorrencia de varias "
        "pessoas. Trabalhos como PIPA, PIPER e modelos contextuais mostram que face, corpo, cena e relacoes entre "
        "pessoas podem ser combinados para melhorar a identificacao em fotos do cotidiano [13, 14, 15]. O presente "
        "trabalho fica em uma versao mais simples: avalia face, global e fusao linear, sem treinar modelos contextuais."
    )
    doc.add_heading("2.4 Fusao e avaliacao de ranking", level=2)
    doc.add_paragraph(
        "A combinacao de diferentes evidencias pode ocorrer antes ou depois da classificacao. Neste trabalho foi "
        "adotada fusao tardia, combinando scores ja calculados, pois embeddings faciais e globais vivem em espacos "
        "semanticos diferentes [16]. A avaliacao usa metricas de recuperacao de informacao, como Precision@K, "
        "Recall@K e mAP, adequadas a sistemas que retornam listas ranqueadas [17]."
    )

    doc.add_heading("3. Metodologia", level=1)
    doc.add_paragraph(
        "A arquitetura implementada foi organizada em tres modulos independentes: indexacao e busca facial, "
        "indexacao e busca global, e fusao tardia. A separacao dos modulos permite avaliar cada estrategia "
        "isoladamente antes de combinar seus resultados."
    )
    add_table(
        doc,
        ["Etapa", "Procedimento", "Artefato produzido"],
        [
            ["Indexacao facial", "Detectar todos os rostos de cada imagem e gerar um embedding por rosto.", "face_embeddings.npy e face_metadata.csv"],
            ["Busca facial", "Extrair o embedding da consulta e comparar com todos os rostos indexados.", "Ranking de fotos completas por pessoa"],
            ["Indexacao global", "Processar a imagem inteira com ResNet50 pre-treinada.", "global_embeddings.npy e global_metadata.csv"],
            ["Busca global", "Comparar o embedding global da consulta contra o indice.", "Ranking por similaridade visual"],
            ["Fusao", "Combinar score facial e score global por soma ponderada.", "Ranking final com pesos alpha e beta"],
        ],
        [1.45, 3.0, 2.05],
    )
    doc.add_paragraph(
        "A similaridade foi calculada por cosseno. Na trilha facial, uma fotografia pode possuir varios rostos, "
        "portanto o indice possui uma entrada por face detectada. No resultado final, as faces correspondentes sao "
        "agrupadas por image_id para retornar a fotografia inteira. Na fusao, foram testados os pesos 0.9/0.1, "
        "0.7/0.3 e 0.5/0.5 para face/global."
    )

    doc.add_heading("4. Experimentos", level=1)
    add_table(
        doc,
        ["Base", "Tamanho usado", "Papel no experimento"],
        [
            ["LFW", "13.233 imagens processadas", "Validacao inicial de embeddings faciais"],
            ["Gallagher", "589 imagens, 931 faces anotadas, 32 identidades", "Busca por pessoa, analise de erros e fusao"],
            ["INRIA Holidays", "1.491 imagens, 500 consultas", "CBIR global por similaridade visual"],
        ],
        [1.45, 2.2, 2.85],
    )
    doc.add_paragraph(
        "O Gallagher foi usado por se aproximar do cenario de albuns pessoais e eventos, com multiplas pessoas "
        "por fotografia. O INRIA Holidays foi usado porque e um benchmark proprio de recuperacao visual, organizado "
        "por grupos de imagens similares. Essa separacao evita comparar tarefas diferentes como se fossem uma unica "
        "avaliacao."
    )

    doc.add_heading("5. Resultados", level=1)
    add_table(
        doc,
        ["Base", "Metodo", "Queries", "P@5", "P@10", "R@10", "mAP", "Tempo"],
        metric_rows(metrics),
        [1.15, 1.45, 0.65, 0.65, 0.65, 0.65, 0.75, 0.75],
    )
    doc.add_paragraph(
        "No Gallagher, a busca facial foi superior ao descritor global e a fusao simples em Precision@5, "
        "Precision@10 e mAP. Isso e esperado, pois a relevancia foi definida pela identidade da pessoa. A ResNet50 "
        "global obteve resultado baixo nesse cenario porque seu vetor representa a imagem inteira, incluindo fundo, "
        "objetos, roupas, iluminacao e outras pessoas. A fusao 0.9/0.1 reduziu a perda em relacao aos pesos mais "
        "equilibrados, mas ainda ficou abaixo da busca facial isolada."
    )
    doc.add_paragraph(
        "No INRIA Holidays, o modulo global apresentou mAP de "
        f"{num(metrics['holidays']['mAP'])}, indicando que a mesma ResNet50 e adequada quando a tarefa e recuperar imagens "
        "visualmente similares. O P@10 e numericamente menor porque muitos grupos possuem poucos itens relevantes, "
        "mas o Recall@10 elevado mostra que as imagens relevantes sao frequentemente recuperadas nos primeiros resultados."
    )
    add_figure(
        doc,
        ROOT / "outputs" / "reports" / "gallagher_face" / "visual_examples" / "face_examples.png",
        "Figura 1 - Busca facial em acervo com multiplas pessoas: o sistema indexa rostos, mas retorna a foto completa.",
    )
    add_figure(
        doc,
        ROOT / "outputs" / "reports" / "holidays_global" / "visual_examples" / "global_examples.png",
        "Figura 2 - Busca global por imagem inteira no INRIA Holidays.",
    )

    doc.add_heading("6. Discussao", level=1)
    doc.add_paragraph(
        "Os resultados sustentam a hipotese de que as duas trilhas sao complementares em intencao de uso, mas nao "
        "necessariamente complementares em qualquer metrica. Quando o usuario procura uma pessoa especifica, o rosto "
        "e a pista dominante. Quando o usuario procura uma cena, objeto ou composicao visual semelhante, o descritor "
        "global se torna mais adequado. Portanto, a interface e a metodologia devem preservar modos de busca separados."
    )
    doc.add_paragraph(
        "O resultado negativo da fusao simples e cientificamente util. Ele mostra que combinar descritores de natureza "
        "diferente sem calibracao pode prejudicar o ranking. Em trabalhos futuros, a fusao pode ser melhorada com "
        "normalizacao de scores, pesos aprendidos, re-ranking, CLIP ou modelos contextuais treinados para albuns "
        "pessoais."
    )

    doc.add_heading("7. Limitacoes", level=1)
    add_bullets(
        doc,
        [
            "A avaliacao facial no Gallagher ainda usa numero limitado de consultas.",
            "O sistema nao usa FAISS; a busca v1 usa NumPy e e adequada para bases pequenas e medias.",
            "O modulo global usa ResNet50 como baseline, sem CLIP ou fine-tuning especifico.",
            "A interface ainda nao permite selecao manual de qual rosto usar quando a consulta possui varias faces.",
            "Nao foi feita avaliacao de vies demografico, privacidade de embeddings ou conformidade juridica completa com dados biometricos.",
        ],
    )

    doc.add_heading("8. Conclusao", level=1)
    doc.add_paragraph(
        "Este trabalho apresentou uma v1 reprodutivel de recuperacao fotografica por conteudo, com busca facial, "
        "busca global e fusao tardia. A implementacao mostrou que embeddings faciais sao mais adequados para busca "
        "por pessoa, enquanto descritores globais ResNet50 sao adequados para recuperacao de imagens visualmente "
        "similares. A fusao linear simples nao melhorou a busca por identidade no Gallagher, indicando que a "
        "combinacao de pistas deve ser calibrada conforme a tarefa. Como continuidade, recomenda-se ampliar a base "
        "de avaliacao, permitir selecao manual do rosto de consulta, testar CLIP, adicionar FAISS e explorar fusao "
        "aprendida."
    )
    add_references(doc)

    output = OUT_DIR / "Artigo_CIBIR_Rascunho.docx"
    doc.save(output)
    return output


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if TMP_IMG_DIR.exists():
        shutil.rmtree(TMP_IMG_DIR)
    metrics = load_metrics()
    report = build_report(metrics)
    article = build_article(metrics)
    print(f"[OK] {report}")
    print(f"[OK] {article}")


if __name__ == "__main__":
    main()
